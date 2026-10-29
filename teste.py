#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Aplica o módulo 'Conciliação Bancária' ao ConvertHub (pywebview)
e torna a barra lateral expansível/recolhível.

Uso: python patch_converthub_conciliacao.py caminho/do/converthub.py
Saída: <nome>_conciliacao.py (novo arquivo) + backup do original.
"""

import sys
import shutil
from datetime import datetime
from pathlib import Path

# ══════════════════════════════════════════════════════════════════════
#  BACKEND (inserido antes de "class Api:")
# ══════════════════════════════════════════════════════════════════════
BACKEND = r'''
# ══════════════════════════════════════════════════════════════════════
#  Conciliação Bancária (extrato OFX × planilha contábil)
# ══════════════════════════════════════════════════════════════════════
_CC_RE_BLOQ = re.compile(
    r"BLOQ|BLOQUEADO|CHEQUE\s*BLOQ|DEP\.?\s*CHEQUE\s*BLOQ", re.IGNORECASE)
_CC_RE_LIBER = re.compile(r"LIBER", re.IGNORECASE)


def _cc_bloqueado(desc):
    """True para CHEQUE BLOQ / BLOQUEADO. False se for liberação (LIBER*)."""
    if not desc:
        return False
    if _CC_RE_LIBER.search(desc):
        return False
    return bool(_CC_RE_BLOQ.search(desc))


def _cc_encoding(filepath):
    try:
        with open(filepath, "rb") as f:
            header = f.read(2000).decode("ascii", errors="ignore")
        m = re.search(r"CHARSET:(\d+)", header, re.I)
        if m:
            cp = f"cp{m.group(1)}"
            "x".encode(cp)
            return cp
    except Exception:
        pass
    return "utf-8"


def _cc_amount(raw):
    if not raw:
        return 0.0
    raw = raw.strip()
    if not raw:
        return 0.0
    if "." in raw and "," in raw:
        raw = raw.replace(".", "").replace(",", ".")
    elif "," in raw:
        raw = raw.replace(",", ".")
    try:
        return float(raw)
    except (ValueError, TypeError):
        return 0.0


def _cc_field(block, tag):
    m = re.search(rf"<{tag}>(.*?)(?:</{tag}>|<|\n|\r)", block, re.I | re.DOTALL)
    return m.group(1).strip() if m else ""


def _cc_date(raw):
    clean = re.sub(r"\[.*?\]", "", raw).strip()
    try:
        return datetime.strptime(clean[:8], "%Y%m%d").strftime("%d/%m/%Y")
    except (ValueError, IndexError):
        return raw


def cc_parse_ofx(filepath):
    enc = _cc_encoding(filepath)
    with open(filepath, "r", encoding=enc, errors="replace") as f:
        content = f.read()
    blocks = re.findall(
        r"<STMTTRN>(.*?)(?=</STMTTRN>|<STMTTRN>|</BANKTRANLIST>|\Z)",
        content, re.DOTALL | re.IGNORECASE)
    out = []
    for block in blocks:
        trnamt = _cc_field(block, "TRNAMT")
        if not trnamt:
            continue
        memo = _cc_field(block, "MEMO") or _cc_field(block, "NAME")
        if _cc_bloqueado(memo):
            continue
        dt = _cc_field(block, "DTPOSTED")
        out.append({
            "Data": _cc_date(dt) if dt else "",
            "Tipo": _cc_field(block, "TRNTYPE"),
            "Descrição": memo,
            "Valor": _cc_amount(trnamt),
        })
    if not out:
        raise ValueError("Nenhum lançamento encontrado no OFX.")
    return out


def cc_parse_ofx_saldo(filepath):
    enc = _cc_encoding(filepath)
    with open(filepath, "r", encoding=enc, errors="replace") as f:
        content = f.read()
    balamt = None
    m = re.search(r"<LEDGERBAL>.*?<BALAMT>(.*?)(?:</BALAMT>|<|\n|\r)",
                  content, re.I | re.DOTALL)
    if not m:
        m = re.search(r"<BALAMT>(.*?)(?:</BALAMT>|<|\n|\r)", content, re.I)
    if m:
        balamt = _cc_amount(m.group(1))
    return {"saldo_final": balamt}


def _cc_norm_date(raw):
    if isinstance(raw, datetime):
        return raw.strftime("%d/%m/%Y")
    if isinstance(raw, date):
        return raw.strftime("%d/%m/%Y")
    s = str(raw).strip()
    for fmt in ("%d/%m/%Y", "%d/%m/%y", "%Y-%m-%d", "%Y%m%d", "%d-%m-%Y"):
        try:
            return datetime.strptime(s, fmt).strftime("%d/%m/%Y")
        except ValueError:
            continue
    return s


def _cc_norm_valor(raw):
    if isinstance(raw, (int, float)):
        return float(raw)
    s = str(raw).strip()
    sinal = 1
    if s.upper().endswith("D"):
        s, sinal = s[:-1].strip(), -1
    elif s.upper().endswith("C"):
        s = s[:-1].strip()
    s = s.replace(".", "").replace(",", ".")
    try:
        return float(s) * sinal
    except ValueError:
        return 0.0


def _cc_hmap(valores):
    hmap = {}
    for ci, v in enumerate(valores):
        v = str(v).strip().lower() if v else ""
        if v == "data":
            hmap["data"] = ci
        elif v == "valor":
            hmap["valor"] = ci
        elif v in ("descrição", "descricao"):
            hmap["desc"] = ci
    if "data" not in hmap or "valor" not in hmap:
        raise ValueError("A planilha precisa ter as colunas 'Data' e 'Valor' "
                         "(e opcionalmente 'Descrição') na primeira linha.")
    return hmap


def _cc_xlsx(filepath):
    import openpyxl
    wb = openpyxl.load_workbook(filepath, data_only=True)
    ws = wb.active
    hmap = _cc_hmap([c.value for c in ws[1]])
    out = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        rv = row[hmap["valor"]] if hmap["valor"] < len(row) else None
        if rv is None or str(rv).strip() == "":
            continue
        desc = ""
        if "desc" in hmap and hmap["desc"] < len(row) and row[hmap["desc"]] is not None:
            desc = str(row[hmap["desc"]]).strip()
        out.append({"Data": _cc_norm_date(row[hmap["data"]]),
                    "Descrição": desc, "Valor": _cc_norm_valor(rv)})
    return out


def _cc_xls(filepath):
    import xlrd
    wb = xlrd.open_workbook(filepath)
    ws = wb.sheet_by_index(0)
    hmap = _cc_hmap([ws.cell_value(0, ci) for ci in range(ws.ncols)])
    out = []
    for ri in range(1, ws.nrows):
        rv = ws.cell_value(ri, hmap["valor"])
        if rv == "" or rv is None:
            continue
        if ws.cell_type(ri, hmap["data"]) == xlrd.XL_CELL_DATE:
            dt = _cc_norm_date(xlrd.xldate_as_datetime(
                ws.cell_value(ri, hmap["data"]), wb.datemode))
        else:
            dt = _cc_norm_date(ws.cell_value(ri, hmap["data"]))
        desc = str(ws.cell_value(ri, hmap["desc"])).strip() if "desc" in hmap else ""
        out.append({"Data": dt, "Descrição": desc, "Valor": _cc_norm_valor(rv)})
    return out


def cc_parse_xls(filepath):
    ext = Path(filepath).suffix.lower()
    ordem = [_cc_xlsx, _cc_xls] if ext in (".xlsx", ".xlsm") else [_cc_xls, _cc_xlsx]
    primeiro = ultimo = None
    for fn in ordem:
        try:
            out = fn(filepath)
            if not out:
                raise ValueError("Nenhum lançamento encontrado na planilha.")
            return out
        except ImportError:
            ultimo = ImportError("Instale: pip install xlrd openpyxl")
        except Exception as exc:
            if primeiro is None:
                primeiro = exc
            ultimo = exc
    raise primeiro or ultimo


# ── motor de cruzamento ────────────────────────────────────────────
# Por dia: 1) match exato (mesmo valor, sinais opostos); 2) combos nos
# restantes (âncora banco, âncora contábil, N:M pequeno). Item usado
# nunca é reaproveitado.
_CC_MAX_ITENS_DP = 400
_CC_MAX_ESTADOS_DP = 10_000
_CC_LIMITE_NM = 8


def _cc_cents(v):
    return round(v * 100)


def _cc_opostos(v1, v2, tol):
    if abs(v1) < tol and abs(v2) < tol:
        return True
    if abs(abs(v1) - abs(v2)) >= tol:
        return False
    return (v1 > 0 and v2 < 0) or (v1 < 0 and v2 > 0)


def _cc_exato(o_list, x_list, tol):
    exatos = []
    x_by_val = defaultdict(list)
    for j, x in enumerate(x_list):
        x_by_val[_cc_cents(abs(x["Valor"]))].append(j)
    usados_x, usados_o = set(), set()
    for i, o in enumerate(o_list):
        for j in x_by_val.get(_cc_cents(abs(o["Valor"])), []):
            if j in usados_x:
                continue
            if _cc_opostos(o["Valor"], x_list[j]["Valor"], tol):
                exatos.append((o, x_list[j]))
                usados_x.add(j)
                usados_o.add(i)
                break
    o_rest = [o for i, o in enumerate(o_list) if i not in usados_o]
    x_rest = [x for j, x in enumerate(x_list) if j not in usados_x]
    return exatos, o_rest, x_rest


def _cc_subset_mask(vals_c, alvo_c, tol_c):
    """Subset-sum aproximado (DP podado por proximidade ao alvo)."""
    n = len(vals_c)
    if n == 0:
        return None
    if n > _CC_MAX_ITENS_DP:
        vals_c = vals_c[:_CC_MAX_ITENS_DP]
    dp = {0: 0}
    melhor_mask, melhor_dist = None, None
    for i, c in enumerate(vals_c):
        novos = {}
        bit = 1 << i
        for s, mask in dp.items():
            ns = s + c
            if ns not in dp and ns not in novos:
                novos[ns] = mask | bit
        dp.update(novos)
        for s, mask in novos.items():
            dist = abs(s - alvo_c)
            if dist <= tol_c and (melhor_dist is None or dist < melhor_dist):
                melhor_mask, melhor_dist = mask, dist
        if melhor_dist is not None:
            return melhor_mask
        if len(dp) > _CC_MAX_ESTADOS_DP:
            itens = sorted(dp.items(), key=lambda kv: abs(kv[0] - alvo_c))
            dp = dict(itens[:_CC_MAX_ESTADOS_DP])
    return melhor_mask


def _cc_por_ancora(ancoras, pool, tol, ancora_e_banco):
    combos = []
    tol_c = max(1, round(tol * 100))
    anc_idx = list(range(len(ancoras)))
    pool_idx = list(range(len(pool)))
    for i in list(anc_idx):
        if not pool_idx:
            break
        alvo_c = -_cc_cents(ancoras[i]["Valor"])
        vals_c = [_cc_cents(pool[j]["Valor"]) for j in pool_idx]
        mask = _cc_subset_mask(vals_c, alvo_c, tol_c)
        if mask is None:
            continue
        usados = [pool_idx[k] for k in range(len(pool_idx)) if mask & (1 << k)]
        if not usados:
            continue
        grupo = [pool[j] for j in usados]
        combos.append(([ancoras[i]], grupo) if ancora_e_banco
                      else (grupo, [ancoras[i]]))
        anc_idx.remove(i)
        for j in usados:
            pool_idx.remove(j)
    return combos, [ancoras[i] for i in anc_idx], [pool[j] for j in pool_idx]


def _cc_nm(o_rest, x_rest, tol):
    if not o_rest or not x_rest:
        return [], o_rest, x_rest
    if len(o_rest) > _CC_LIMITE_NM or len(x_rest) > _CC_LIMITE_NM:
        return [], o_rest, x_rest
    tol_c = max(1, round(tol * 100))
    o_idx, x_idx, combos = list(range(len(o_rest))), list(range(len(x_rest))), []
    progrediu = True
    while progrediu and o_idx and x_idx:
        progrediu = False
        achou = None
        n_o, n_x = len(o_idx), len(x_idx)
        for total in range(2, n_o + n_x + 1):
            for tam_o in range(1, min(total - 1, n_o) + 1):
                tam_x = total - tam_o
                if tam_x < 1 or tam_x > n_x:
                    continue
                for oc in combinations(o_idx, tam_o):
                    so = sum(_cc_cents(o_rest[i]["Valor"]) for i in oc)
                    for xc in combinations(x_idx, tam_x):
                        sx = sum(_cc_cents(x_rest[j]["Valor"]) for j in xc)
                        if abs(sx + so) <= tol_c:
                            achou = (oc, xc)
                            break
                    if achou:
                        break
                if achou:
                    break
            if achou:
                break
        if achou:
            oc, xc = achou
            combos.append(([o_rest[i] for i in oc], [x_rest[j] for j in xc]))
            o_idx = [i for i in o_idx if i not in oc]
            x_idx = [j for j in x_idx if j not in xc]
            progrediu = True
    return combos, [o_rest[i] for i in o_idx], [x_rest[j] for j in x_idx]


def _cc_combinar(o_rest, x_rest, tol):
    combos = []
    c1, o_rest, x_rest = _cc_por_ancora(o_rest, x_rest, tol, True)
    combos += c1
    c2, x_rest, o_rest = _cc_por_ancora(x_rest, o_rest, tol, False)
    combos += c2
    c3, o_rest, x_rest = _cc_nm(o_rest, x_rest, tol)
    combos += c3
    return combos, o_rest, x_rest


def _cc_ordem_data(d):
    try:
        return datetime.strptime(d, "%d/%m/%Y")
    except ValueError:
        return datetime.max


def cc_cruzar(ofx_data, xls_data, tolerancia=0.01,
              saldo_final_ofx=None, saldo_anterior=None):
    ofx_dia, xls_dia = defaultdict(list), defaultdict(list)
    for it in ofx_data:
        ofx_dia[it["Data"]].append(dict(it))
    for it in xls_data:
        xls_dia[it["Data"]].append(dict(it))
    dias = sorted(set(ofx_dia) | set(xls_dia), key=_cc_ordem_data)

    mov_periodo = round(sum(i["Valor"] for i in ofx_data), 2)
    saldo_inicial = (round(saldo_final_ofx - mov_periodo, 2)
                     if saldo_final_ofx is not None else 0.0)
    ant = round(float(saldo_anterior), 2) if saldo_anterior is not None else 0.0

    resultado = {}
    saldo_ofx = saldo_xls = saldo_inicial
    for dia in dias:
        o_list, x_list = ofx_dia.get(dia, []), xls_dia.get(dia, [])
        exatos, o_rest, x_rest = _cc_exato(o_list, x_list, tolerancia)
        combos = []
        if o_rest and x_rest:
            combos, o_rest, x_rest = _cc_combinar(o_rest, x_rest, tolerancia)
        t_o = round(sum(i["Valor"] for i in o_list), 2)
        t_x = round(sum(i["Valor"] for i in x_list), 2)
        saldo_ofx = round(saldo_ofx + t_o, 2)
        saldo_xls = round(saldo_xls + (-t_x), 2)
        resultado[dia] = {
            "total_ofx": t_o, "total_xls": t_x,
            "diferenca": round(t_o + t_x, 2),
            "saldo_ofx": saldo_ofx, "saldo_xls": saldo_xls,
            "cruzados_exatos": exatos, "cruzados_combo": combos,
            "so_ofx": o_rest, "so_xls": x_rest,
        }

    final_ofx = round(saldo_ofx + ant, 2)
    final_xls = round(saldo_xls + ant, 2)
    if dias:
        resultado[dias[-1]]["saldo_ofx"] = final_ofx
        resultado[dias[-1]]["saldo_xls"] = final_xls
    meta = {
        "saldo_final_ofx": final_ofx, "saldo_final_xls": final_xls,
        "saldo_oficial_ofx": saldo_final_ofx,
        "ultimo_dia": dias[-1] if dias else None,
    }
    return resultado, meta


def _cc_ler_saldo(raw):
    raw = (raw or "").strip()
    if not raw:
        return None
    s = raw.replace("R$", "").replace(" ", "")
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".")
    elif "," in s:
        s = s.replace(",", ".")
    try:
        return float(s)
    except ValueError:
        raise ValueError(f"Saldo anterior inválido: '{raw}'. "
                         "Use algo como 1234,56 ou -1.234,56")


def _cc_payload(res, meta, ant):
    def it(i):
        return {"desc": i.get("Descrição", ""), "valor": i["Valor"]}

    dias = []
    tot_ex = tot_co = tot_b = tot_c = dias_dif = 0
    tot_dif = 0.0
    ultimo = meta.get("ultimo_dia")
    for dia, d in res.items():
        tot_ex += len(d["cruzados_exatos"])
        tot_co += len(d["cruzados_combo"])
        tot_b += len(d["so_ofx"])
        tot_c += len(d["so_xls"])
        tot_dif += abs(d["diferenca"])
        if abs(d["diferenca"]) >= 0.01:
            dias_dif += 1
        dias.append({
            "dia": dia, "ultimo": dia == ultimo,
            "total_ofx": d["total_ofx"], "total_xls": d["total_xls"],
            "dif": d["diferenca"],
            "saldo_ofx": d["saldo_ofx"], "saldo_xls": d["saldo_xls"],
            "exatos": [{"b": o["Descrição"], "c": x["Descrição"],
                        "vb": o["Valor"], "vc": x["Valor"]}
                       for o, x in d["cruzados_exatos"]],
            "combos": [{"banco": [it(i) for i in ol],
                        "contab": [it(i) for i in xl]}
                       for ol, xl in d["cruzados_combo"]],
            "so_ofx": [it(i) for i in d["so_ofx"]],
            "so_xls": [it(i) for i in d["so_xls"]],
        })
    return {
        "dias": dias, "tot_exatos": tot_ex, "tot_combos": tot_co,
        "tot_banco": tot_b, "tot_contab": tot_c,
        "tot_dif": round(tot_dif, 2), "dias_dif": dias_dif,
        "saldo_final_ofx": meta["saldo_final_ofx"],
        "saldo_final_xls": meta["saldo_final_xls"],
        "saldo_oficial": meta["saldo_oficial_ofx"],
        "ultimo_dia": ultimo, "ant": ant,
    }

'''

# ══════════════════════════════════════════════════════════════════════
#  MÉTODOS DA API (inseridos antes de "def init_info")
# ══════════════════════════════════════════════════════════════════════
API_METHODS = r'''    # ══════════════════════════════════════════════════════════════
    #  Conciliação Bancária
    # ══════════════════════════════════════════════════════════════
    def _cc_s(self):
        if not hasattr(self, "_cc"):
            self._cc = {"ofx": [], "xls": [], "meta": {}}
        return self._cc

    def conc_pick(self, kind):
        if kind == "ofx":
            ft = ("OFX (*.ofx)", "Todos (*.*)")
        else:
            ft = ("Planilhas (*.xls;*.xlsx)", "Todos (*.*)")
        r = self._dialog("OPEN", file_types=ft)
        return r[0] if r else None

    def conc_carregar(self, kind, path):
        s = self._cc_s()
        p = Path(path or "")
        if not p.is_file():
            return {"ok": False, "erro": "Arquivo não encontrado."}
        try:
            if kind == "ofx":
                s["ofx"] = cc_parse_ofx(str(p))
                s["meta"] = cc_parse_ofx_saldo(str(p))
                return {"ok": True, "nome": p.name, "n": len(s["ofx"]),
                        "saldo": s["meta"].get("saldo_final")}
            s["xls"] = cc_parse_xls(str(p))
            return {"ok": True, "nome": p.name, "n": len(s["xls"]),
                    "saldo": None}
        except Exception as exc:
            if kind == "ofx":
                s["ofx"], s["meta"] = [], {}
            else:
                s["xls"] = []
            return {"ok": False, "erro": str(exc)}

    def conc_conciliar(self, saldo_txt):
        s = self._cc_s()
        if not s["ofx"] or not s["xls"]:
            return {"ok": False, "erro": "Carregue o OFX e a planilha contábil."}
        try:
            ant = _cc_ler_saldo(saldo_txt)
        except ValueError as exc:
            return {"ok": False, "erro": str(exc)}
        try:
            res, meta = cc_cruzar(
                s["ofx"], s["xls"],
                saldo_final_ofx=s["meta"].get("saldo_final"),
                saldo_anterior=ant)
            return {"ok": True, **_cc_payload(res, meta, ant)}
        except Exception as exc:
            return {"ok": False, "erro": f"{type(exc).__name__}: {exc}",
                    "trace": traceback.format_exc()}

'''

# ══════════════════════════════════════════════════════════════════════
#  CSS
# ══════════════════════════════════════════════════════════════════════
CSS = r'''
/* ===== sidebar expansível ===== */
aside { transition:width .18s ease; overflow:hidden; }
aside.collapsed { width:68px; }
aside.collapsed nav p { display:none; }
aside.collapsed .sidebar-item { font-size:0; justify-content:center; padding-left:0; padding-right:0; }
aside.collapsed .brand-text { display:none; }
aside.collapsed .side-foot { display:none; }
aside.collapsed #theme-btn { font-size:0; }
aside.collapsed #theme-btn::before { content:"\25D0"; font-size:16px; }
aside.collapsed > div:first-child { padding-left:0; padding-right:0; }
aside.collapsed > div:first-child > div { justify-content:center; }
#sb-toggle.btn-blue { background:transparent !important; color:#64748b !important; border-color:#cbd5e1 !important; }
#sb-toggle.btn-blue:hover { background:#f1f5f9 !important; }
html.dark #sb-toggle.btn-blue { color:#b8b8b8 !important; border-color:#2a2a2a !important; }
html.dark #sb-toggle.btn-blue:hover { background:#1a1a1a !important; }

/* ===== conciliação ===== */
.cc-wrap table { border-collapse:separate; border-spacing:0; }
.cc-wrap thead th {
  position:sticky; top:0; z-index:5; background:#ffffff; color:#334155;
  box-shadow:inset 0 -2px 0 #0284c7; padding:8px; text-align:left;
  font-size:11px; text-transform:uppercase; letter-spacing:.03em;
}
.cc-wrap td { padding:6px 8px; vertical-align:middle; word-break:break-word; overflow-wrap:anywhere; line-height:1.3; }
.cc-wrap td.num { text-align:right; font-variant-numeric:tabular-nums; white-space:nowrap; }
.cc-day { cursor:pointer; font-weight:600; }
.cc-day:hover td { background:#e6f2fd; }
.cc-day .cc-arrow { display:inline-block; width:14px; font-size:10px; color:#64748b; }
html.dark .cc-wrap thead th { background:#111111; color:#cfe8fa; }
html.dark .cc-day:hover td { background:#10202c; }
'''

# ══════════════════════════════════════════════════════════════════════
#  HTML
# ══════════════════════════════════════════════════════════════════════
SIDEBAR_BTN = r'''<button data-module="conciliacao" class="sidebar-item w-full flex items-center gap-3 px-3 py-2.5 rounded-lg text-sm font-medium text-slate-500 hover:bg-slate-50 hover:text-slate-700 transition-colors text-left">
        <svg class="w-5 h-5 shrink-0" fill="none" stroke="currentColor" viewBox="0 0 24 24"><path stroke-linecap="round" stroke-linejoin="round" stroke-width="1.75" d="M9 5H7a2 2 0 00-2 2v12a2 2 0 002 2h10a2 2 0 002-2V7a2 2 0 00-2-2h-2M9 5a2 2 0 002 2h2a2 2 0 002-2M9 5a2 2 0 012-2h2a2 2 0 012 2m-6 9l2 2 4-4"/></svg>
        Conciliação
      </button>
      '''

THEME_BTN_NEW = r'''<button id="theme-btn" class="w-full px-3 py-2 text-sm">Tema: Claro</button>
      <button id="sb-toggle" class="w-full px-3 py-2 text-sm mt-2" title="Recolher menu">« Recolher</button>'''

HOME_CARD_TAIL = r'''PDF ou OFX, perfis, memória de códigos e regras de palavras-chave</p>
          </button>
          <button data-goto="conciliacao" class="goto-btn group bg-white rounded-xl border border-slate-200 p-5 shadow-sm hover:border-primary-300 hover:shadow-md transition-all text-left">
            <div class="w-10 h-10 rounded-lg bg-primary-50 text-primary-600 flex items-center justify-center mb-3 group-hover:bg-primary-100 transition-colors">
              <svg class="w-5 h-5" fill="none" stroke="currentColor" viewBox="0 0 24 24"><path stroke-linecap="round" stroke-linejoin="round" stroke-width="1.75" d="M9 5H7a2 2 0 00-2 2v12a2 2 0 002 2h10a2 2 0 002-2V7a2 2 0 00-2-2h-2M9 5a2 2 0 002 2h2a2 2 0 002-2M9 5a2 2 0 012-2h2a2 2 0 012 2m-6 9l2 2 4-4"/></svg>
            </div>
            <p class="font-medium text-slate-800">Conciliação Bancária</p>
            <p class="text-xs text-slate-400 mt-1">Cruza o OFX do banco com a planilha contábil: matches exatos, combinações e saldo por dia</p>'''

MODULE_HTML = r'''<!-- ========== CONCILIAÇÃO ========== -->
    <div id="module-conciliacao" class="module-content hidden">
      <div class="flex flex-col" style="height:100vh">
        <header class="bg-white border-b border-slate-200 px-6 py-3 shrink-0">
          <div class="flex items-center justify-between gap-3 flex-wrap">
            <div>
              <h1 class="text-lg font-semibold text-slate-900">Conciliação Bancária</h1>
              <p class="text-xs text-slate-500">Cruze o extrato OFX com a planilha contábil — saldo por dia e saldo final no último dia</p>
            </div>
            <div class="flex gap-2">
              <button id="cc-run" disabled class="px-4 py-2 text-sm">Conciliar</button>
              <button id="cc-exp" disabled class="px-3 py-2 text-sm">Expandir tudo</button>
              <button id="cc-col" disabled class="px-3 py-2 text-sm">Recolher tudo</button>
            </div>
          </div>
          <div class="flex items-stretch gap-3 mt-3 flex-wrap">
            <div class="flex-1 flex items-center gap-2 rounded-lg border border-slate-200 px-3 py-1.5" style="min-width:240px">
              <span class="text-xs font-semibold text-amber-600 shrink-0">OFX (Banco)</span>
              <span id="cc-ofx-lbl" class="flex-1 min-w-0 truncate text-xs text-slate-500">Nenhum arquivo</span>
              <button id="cc-ofx-btn" class="px-3 py-1 text-xs">Procurar</button>
            </div>
            <div class="flex-1 flex items-center gap-2 rounded-lg border border-slate-200 px-3 py-1.5" style="min-width:240px">
              <span class="text-xs font-semibold text-amber-600 shrink-0">Planilha (.xls/.xlsx)</span>
              <span id="cc-xls-lbl" class="flex-1 min-w-0 truncate text-xs text-slate-500">Nenhum arquivo</span>
              <button id="cc-xls-btn" class="px-3 py-1 text-xs">Procurar</button>
            </div>
            <div class="flex items-center gap-2 rounded-lg border border-slate-200 px-3 py-1.5">
              <span class="text-xs font-semibold text-slate-600 shrink-0">Saldo anterior</span>
              <input id="cc-ant" class="w-28 rounded border border-slate-200 px-2 py-1 text-sm focus:outline-none focus:ring-2 focus:ring-primary-500" placeholder="0,00">
            </div>
          </div>
          <p id="cc-status" class="text-xs text-slate-500 mt-2">Carregue o OFX e a planilha para conciliar.</p>
          <div id="cc-prog" class="prog mt-2 hidden"><i></i></div>
        </header>

        <div class="px-6 pt-3 shrink-0">
          <div style="display:grid;grid-template-columns:repeat(auto-fit,minmax(125px,1fr));gap:8px">
            <div class="bg-white rounded-xl border border-slate-200 shadow-sm px-3 py-2 text-center"><p class="text-xs text-slate-400">Exatos</p><p id="cc-c-ex" class="text-base font-semibold text-emerald-600">—</p></div>
            <div class="bg-white rounded-xl border border-slate-200 shadow-sm px-3 py-2 text-center"><p class="text-xs text-slate-400">Combos</p><p id="cc-c-co" class="text-base font-semibold text-amber-600">—</p></div>
            <div class="bg-white rounded-xl border border-slate-200 shadow-sm px-3 py-2 text-center"><p class="text-xs text-slate-400">Só no Banco</p><p id="cc-c-b" class="text-base font-semibold text-red-500">—</p></div>
            <div class="bg-white rounded-xl border border-slate-200 shadow-sm px-3 py-2 text-center"><p class="text-xs text-slate-400">Só na Contab.</p><p id="cc-c-c" class="text-base font-semibold text-red-500">—</p></div>
            <div class="bg-white rounded-xl border border-slate-200 shadow-sm px-3 py-2 text-center"><p class="text-xs text-slate-400">Diferença total</p><p id="cc-c-dif" class="text-base font-semibold text-sky-700">—</p></div>
            <div class="bg-white rounded-xl border border-slate-200 shadow-sm px-3 py-2 text-center"><p class="text-xs text-slate-400">Saldo final Banco</p><p id="cc-c-sb" class="text-base font-semibold text-emerald-600">—</p></div>
            <div class="bg-white rounded-xl border border-slate-200 shadow-sm px-3 py-2 text-center"><p class="text-xs text-slate-400">Saldo final Contábil</p><p id="cc-c-sc" class="text-base font-semibold text-amber-600">—</p></div>
          </div>
        </div>

        <div class="px-6 py-3 flex-1 min-h-0">
          <div class="bg-white rounded-xl border border-slate-200 shadow-sm h-full overflow-hidden">
            <div class="cc-wrap h-full overflow-auto">
              <table class="w-full text-xs" style="table-layout:fixed">
                <colgroup><col style="width:9%"><col style="width:26%"><col style="width:12%"><col style="width:12%"><col style="width:14%"><col style="width:14%"><col style="width:13%"></colgroup>
                <thead><tr>
                  <th>Data</th><th>Descrição</th>
                  <th style="text-align:right">Mov. Banco</th><th style="text-align:right">Mov. Contab.</th>
                  <th style="text-align:right">Saldo Banco</th><th style="text-align:right">Saldo Contábil</th>
                  <th style="text-align:center">Status</th>
                </tr></thead>
                <tbody id="cc-body"></tbody>
              </table>
              <div id="cc-empty" class="p-10 text-center text-sm text-slate-400">Carregue os arquivos e clique em Conciliar.</div>
            </div>
          </div>
        </div>
      </div>
    </div>

    '''

# ══════════════════════════════════════════════════════════════════════
#  JS
# ══════════════════════════════════════════════════════════════════════
JS = r'''
<script>
/* ===== sidebar expansível ===== */
(function () {
  var aside = document.querySelector('aside');
  var btn = document.getElementById('sb-toggle');
  if (!aside || !btn) return;
  document.querySelectorAll('.sidebar-item').forEach(function (b) {
    b.title = b.textContent.replace(/\s+/g, ' ').trim();
  });
  function setCollapsed(c, save) {
    aside.classList.toggle('collapsed', c);
    btn.textContent = c ? '»' : '« Recolher';
    btn.title = c ? 'Expandir menu' : 'Recolher menu';
    if (save) { try { localStorage.setItem('hub-sb', c ? '1' : '0'); } catch (e) {} }
  }
  var saved = null;
  try { saved = localStorage.getItem('hub-sb'); } catch (e) {}
  setCollapsed(saved === null ? window.innerWidth < 1250 : saved === '1', false);
  btn.addEventListener('click', function () {
    setCollapsed(!aside.classList.contains('collapsed'), true);
  });
})();

/* ===== conciliação ===== */
(function () {
  const C = { ofx: false, xls: false, busy: false };
  const fmt = v => (v === null || v === undefined || v === '') ? '' : brl(v);

  function ready() {
    $('#cc-run').disabled = C.busy || !(C.ofx && C.xls);
  }
  function status(t, cls) {
    const s = $('#cc-status');
    s.textContent = t;
    s.className = 'text-xs mt-2 ' + (cls || 'text-slate-500');
  }

  async function pick(kind) {
    const p = await py('conc_pick', kind);
    if (!p) return;
    const lbl = $('#cc-' + kind + '-lbl');
    lbl.textContent = 'Carregando…';
    let r;
    try { r = await py('conc_carregar', kind, p); }
    catch (e) { lbl.textContent = 'Nenhum arquivo'; return; }
    if (!r.ok) {
      C[kind] = false; lbl.textContent = 'Nenhum arquivo'; ready();
      return toast(r.erro, 'err');
    }
    C[kind] = true;
    lbl.textContent = r.nome + '  ·  ' + r.n + ' lanç.';
    lbl.title = p;
    let extra = '';
    if (kind === 'ofx' && r.saldo !== null && r.saldo !== undefined) extra = '  |  Saldo extrato: ' + brl(r.saldo);
    status((kind === 'ofx' ? 'OFX' : 'Planilha') + ' carregado(a): ' + r.n + ' lançamento(s)' + extra,
           'text-emerald-700');
    ready();
  }

  function det(i, desc, vb, vc, st, cls, sub) {
    return `<tr class="cc-det hidden border-t border-slate-100" data-p="${i}"><td></td>
      <td class="${cls}" style="${sub ? 'padding-left:24px;opacity:.85' : ''}">${esc(desc)}</td>
      <td class="num ${cls}">${fmt(vb)}</td><td class="num ${cls}">${fmt(vc)}</td><td></td><td></td>
      <td class="${cls}" style="text-align:center">${esc(st)}</td></tr>`;
  }
  function sec(i, t) {
    return `<tr class="cc-det hidden" data-p="${i}"><td></td><td colspan="6" class="text-xs italic text-slate-400">${esc(t)}</td></tr>`;
  }
  const sum = a => a.reduce((s, x) => s + x.valor, 0);

  function render(r) {
    const rows = [];
    r.dias.forEach((d, i) => {
      const ok = Math.abs(d.dif) < 0.01;
      const cls = d.ultimo ? 'text-sky-700' : (ok ? 'text-emerald-600' : 'text-red-500');
      let st = d.ultimo ? '★ Saldo Final' : (ok ? '✓ Conciliado' : '⚠ DIF ' + brl(d.dif));
      if (d.ultimo && !ok) st += ' · DIF ' + brl(d.dif);
      const partes = [];
      if (d.exatos.length) partes.push(d.exatos.length + ' exato(s)');
      if (d.combos.length) partes.push(d.combos.length + ' combo(s)');
      if (d.so_ofx.length) partes.push(d.so_ofx.length + ' só banco');
      if (d.so_xls.length) partes.push(d.so_xls.length + ' só contab.');
      let desc = partes.join(' · ') || '— nada neste dia —';
      if (d.ultimo) desc = '▶ SALDO FINAL' + (r.ant ? ' (período + ant. ' + brl(r.ant) + ')' : '') + ' · ' + desc;
      rows.push(`<tr class="cc-day bg-slate-50 border-t border-slate-100 ${cls}" data-i="${i}">
        <td><span class="cc-arrow">▶</span>${esc(d.dia)}</td><td>${esc(desc)}</td>
        <td class="num">${fmt(d.total_ofx)}</td><td class="num">${fmt(d.total_xls)}</td>
        <td class="num">${fmt(d.saldo_ofx)}</td><td class="num">${fmt(d.saldo_xls)}</td>
        <td style="text-align:center">${esc(st)}</td></tr>`);
      if (d.exatos.length) {
        rows.push(sec(i, '── Matches exatos (' + d.exatos.length + ') ──'));
        d.exatos.forEach(m => rows.push(det(i, m.b + '  ↔  ' + m.c, m.vb, m.vc, '✓ Exato', 'text-emerald-600')));
      }
      if (d.combos.length) {
        rows.push(sec(i, '── Combinações (' + d.combos.length + ') ──'));
        d.combos.forEach(c => {
          const od = c.banco.map(x => x.desc).join(' + '), xd = c.contab.map(x => x.desc).join(' + ');
          rows.push(det(i, '[' + od + ']  ↔  [' + xd + ']', sum(c.banco), sum(c.contab),
                        '⚡ Combo ' + c.banco.length + ':' + c.contab.length, 'text-amber-600'));
          if (c.banco.length > 1 || c.contab.length > 1) {
            c.banco.forEach(x => rows.push(det(i, 'banco: ' + x.desc, x.valor, null, '', 'text-amber-600', true)));
            c.contab.forEach(x => rows.push(det(i, 'contab.: ' + x.desc, null, x.valor, '', 'text-amber-600', true)));
          }
        });
      }
      if (d.so_ofx.length) {
        rows.push(sec(i, '── Só no Banco (' + d.so_ofx.length + ') ──'));
        d.so_ofx.forEach(x => rows.push(det(i, x.desc, x.valor, null, '⚠ Banco', 'text-red-500')));
      }
      if (d.so_xls.length) {
        rows.push(sec(i, '── Só na Contab. (' + d.so_xls.length + ') ──'));
        d.so_xls.forEach(x => rows.push(det(i, x.desc, null, x.valor, '⚠ Contab.', 'text-red-500')));
      }
    });
    $('#cc-body').innerHTML = rows.join('');
    $('#cc-empty').classList.toggle('hidden', r.dias.length > 0);

    $('#cc-c-ex').textContent = r.tot_exatos;
    $('#cc-c-co').textContent = r.tot_combos;
    $('#cc-c-b').textContent = r.tot_banco;
    $('#cc-c-c').textContent = r.tot_contab;
    $('#cc-c-dif').textContent = brl(r.tot_dif);
    $('#cc-c-sb').textContent = brl(r.saldo_final_ofx);
    $('#cc-c-sc').textContent = brl(r.saldo_final_xls);

    let msg = 'Saldo em ' + (r.ultimo_dia || '—') + ' → Banco: ' + brl(r.saldo_final_ofx) +
              ' | Contábil: ' + brl(r.saldo_final_xls);
    if (r.saldo_oficial !== null && r.saldo_oficial !== undefined) msg += ' | Extrato OFX: ' + brl(r.saldo_oficial);
    if (r.dias_dif === 0) status('Conciliação concluída: todos os dias fecharam · ' + msg, 'text-emerald-700');
    else status('Conciliação concluída: ' + r.dias_dif + ' dia(s) com diferença · ' + msg, 'text-amber-600');
    $('#cc-exp').disabled = false;
    $('#cc-col').disabled = false;
  }

  function toggleDay(tr, open) {
    const i = tr.dataset.i;
    if (open === undefined) open = tr.dataset.open !== '1';
    tr.dataset.open = open ? '1' : '0';
    tr.querySelector('.cc-arrow').textContent = open ? '▼' : '▶';
    $$('#cc-body tr.cc-det[data-p="' + i + '"]').forEach(x => x.classList.toggle('hidden', !open));
  }

  async function run() {
    if (C.busy) return;
    C.busy = true; ready();
    $('#cc-prog').classList.remove('hidden');
    status('Conciliando… isso pode levar alguns instantes.', 'text-amber-600');
    let r;
    try { r = await py('conc_conciliar', $('#cc-ant').value); }
    catch (e) { r = null; }
    finally { C.busy = false; $('#cc-prog').classList.add('hidden'); ready(); }
    if (!r) return status('Falha ao conciliar.', 'text-red-500');
    if (!r.ok) { status(r.erro, 'text-red-500'); return toast(r.erro, 'err'); }
    render(r);
  }

  $('#cc-ofx-btn').onclick = () => pick('ofx');
  $('#cc-xls-btn').onclick = () => pick('xls');
  $('#cc-run').onclick = run;
  $('#cc-ant').onkeydown = e => { if (e.key === 'Enter' && !$('#cc-run').disabled) run(); };
  $('#cc-exp').onclick = () => $$('#cc-body tr.cc-day').forEach(tr => toggleDay(tr, true));
  $('#cc-col').onclick = () => $$('#cc-body tr.cc-day').forEach(tr => toggleDay(tr, false));
  $('#cc-body').addEventListener('click', e => {
    const tr = e.target.closest('tr.cc-day');
    if (tr) toggleDay(tr);
  });
})();
</script>
'''


def trocar(txt, ancora, novo, nome):
    n = txt.count(ancora)
    if n != 1:
        raise ValueError(
            f"Ponto de inserção '{nome}' encontrado {n}x (esperado 1). "
            "O arquivo parece diferente da versão esperada do ConvertHub.")
    return txt.replace(ancora, novo)


def aplicar(txt):
    if "conc_conciliar" in txt or 'id="module-conciliacao"' in txt:
        raise ValueError("O módulo de Conciliação já está presente neste arquivo.")

    txt = trocar(txt, "from pathlib import Path",
                 "from pathlib import Path\nfrom collections import defaultdict\n"
                 "from itertools import combinations", "imports")
    txt = trocar(txt, "\nclass Api:\n",
                 "\n" + BACKEND + "\nclass Api:\n", "class Api")
    txt = trocar(txt, "    def init_info(self):",
                 API_METHODS + "    def init_info(self):", "init_info")
    txt = trocar(txt, "</style>", CSS + "</style>", "css")
    txt = trocar(txt,
                 '<span class="font-semibold text-lg tracking-tight">',
                 '<span class="brand-text font-semibold text-lg tracking-tight">',
                 "logo")
    txt = trocar(txt, '<div class="pt-4 mt-2 border-t border-slate-100">',
                 SIDEBAR_BTN + '<div class="pt-4 mt-2 border-t border-slate-100">',
                 "menu lateral")
    txt = trocar(txt,
                 '<button id="theme-btn" class="w-full px-3 py-2 text-sm">Tema: Claro</button>',
                 THEME_BTN_NEW, "botão de tema")
    txt = trocar(txt,
                 '<div class="px-4 py-4 border-t border-slate-100 text-xs text-slate-400">',
                 '<div class="side-foot px-4 py-4 border-t border-slate-100 text-xs text-slate-400">',
                 "rodapé da sidebar")
    txt = trocar(txt,
                 "PDF ou OFX, perfis, memória de códigos e regras de palavras-chave</p>",
                 HOME_CARD_TAIL, "card da home")
    txt = trocar(txt, "<!-- ========== SISTEMA ========== -->",
                 MODULE_HTML + "<!-- ========== SISTEMA ========== -->",
                 "módulo html")
    txt = trocar(txt, "</body>", JS + "</body>", "javascript")
    return txt


def main():
    if len(sys.argv) != 2:
        print("Uso: python patch_converthub_conciliacao.py <arquivo_converthub.py>")
        sys.exit(1)

    src = Path(sys.argv[1]).expanduser().resolve()
    if not src.is_file():
        print(f"Erro: arquivo não encontrado: {src}")
        sys.exit(1)

    try:
        raw = src.read_bytes()
        try:
            texto = raw.decode("utf-8-sig")
            bom = raw.startswith(b"\xef\xbb\xbf")
        except UnicodeDecodeError:
            texto = raw.decode("cp1252")
            bom = False
        crlf = "\r\n" in texto
        texto = texto.replace("\r\n", "\n")

        if "ConvertHub" not in texto or "class Api:" not in texto:
            raise ValueError("O arquivo informado não parece ser o ConvertHub (pywebview).")

        novo = aplicar(texto)
        compile(novo, str(src), "exec")  # valida sintaxe antes de gravar

        if crlf:
            novo = novo.replace("\n", "\r\n")

        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        backup = src.with_name(f"{src.stem}.backup_{stamp}{src.suffix}")
        shutil.copy2(src, backup)

        destino = src.with_name(f"{src.stem}_conciliacao{src.suffix}")
        if destino.exists():
            destino = src.with_name(f"{src.stem}_conciliacao_{stamp}{src.suffix}")
        destino.write_bytes(
            (b"\xef\xbb\xbf" if bom else b"") + novo.encode("utf-8"))

        print(f"Backup do original: {backup}")
        print(f"Arquivo corrigido:  {destino}")
    except (ValueError, SyntaxError) as exc:
        print(f"Erro: {exc}")
        sys.exit(2)
    except OSError as exc:
        print(f"Erro de arquivo: {exc}")
        sys.exit(3)
    except Exception as exc:
        print(f"Erro inesperado: {type(exc).__name__}: {exc}")
        sys.exit(4)


if __name__ == "__main__":
    main()